#!/usr/bin/env python3
"""
Destatis Municipal Directory (GV3Q) Extractor
============================================
Reads the official Excel file 'AuszugGV3QAktuell.xlsx' from the Federal
Statistical Office of Germany (Destatis) and exports all municipalities and
cities (record type / Satzart 60), enriched with parent state and district,
to JSON and CSV formats.

Standard fields:
- ars: Official Regional Key (Amtlicher Regionalschlüssel, 12 digits)
- state: Name of the federal state (from record type 10)
- district: Name of the district / independent city (from record type 40)
- name: Official municipality name
- slug: URL-safe slug with German umlaut transliteration
- type: Type of municipality ('City', 'Municipality', 'Unincorporated area')
- postal_code: 5-digit postal code of the administrative headquarters
- country_code: ISO 3166-1 alpha-2 country code (always 'DE')
- population: Population (based on 2022 Census)
- area_km2: Area in km²
- latitude: Geographic latitude
- longitude: Geographic longitude

https://www.destatis.de/DE/Themen/Laender-Regionen/Regionales/Gemeindeverzeichnis/_inhalt.html
"""

import argparse
from collections import Counter
import csv
import json
import os
import sys
import time
from typing import Any, Dict, List, Optional
import openpyxl
from slugify import slugify

# Destatis official Textkennzeichen (TKZ) mapping for Satzart 60
TKZ_TO_TYPE: Dict[str, str] = {
    "60": "Municipality",         # Markt (classified as Municipality)
    "61": "City",                 # Kreisfreie Stadt
    "62": "City",                 # Stadtkreis (Baden-Württemberg)
    "63": "City",                 # Kreisangehörige Stadt / Stadt
    "64": "Municipality",         # Kreisangehörige Gemeinde / Gemeinde
    "65": "Unincorporated area",   # Gemeindefreier Bezirk
    "66": "Unincorporated area",   # Gemeindefreies Gebiet
    "67": "City",                 # Große Kreisstadt
}


def determine_type(tkz: Any, raw_name: Optional[str] = None) -> str:
    """Determines whether a municipality is a City, Municipality, or Unincorporated area."""
    if tkz is not None:
        clean_tkz = str(tkz).strip()
        if clean_tkz in TKZ_TO_TYPE:
            return TKZ_TO_TYPE[clean_tkz]

    if raw_name:
        lower_name = raw_name.lower()
        if any(term in lower_name for term in [", stadt", ", hansestadt", ", landeshauptstadt", ", st", ", gkst"]):
            return "City"
        if "gemfr." in lower_name or "gemeindefreies" in lower_name:
            return "Unincorporated area"

    return "Municipality"


GERMAN_SLUG_REPLACEMENTS = [
    ["ä", "ae"],
    ["ö", "oe"],
    ["ü", "ue"],
    ["ß", "ss"],
    ["Ä", "ae"],
    ["Ö", "oe"],
    ["Ü", "ue"],
]


def generate_slug(text: str) -> str:
    """Generates a safe URL slug with German umlaut transliteration."""
    return slugify(text, replacements=GERMAN_SLUG_REPLACEMENTS)


def ensure_unique_slugs_and_names(records: List[Dict[str, Any]]) -> None:
    """
    Ensures that every municipality has a globally unique slug.
    If a slug occurs multiple times, appends the district in parentheses to the name.
    If collisions still remain within the same district, further disambiguates with type/plz.
    """
    # Pass 1: detect duplicate slugs and append district in parentheses to the name
    slug_counts = Counter(r["slug"] for r in records)
    for r in records:
        if slug_counts[r["slug"]] > 1 and r.get("district"):
            if not r["name"].endswith(f"({r['district']})"):
                r["name"] = f"{r['name']} ({r['district']})"
            r["slug"] = generate_slug(r["name"])

    # Pass 2: handle remaining collisions occurring within the same district
    slug_counts2 = Counter(r["slug"] for r in records)
    for r in records:
        if slug_counts2[r["slug"]] > 1:
            if r.get("type") == "Unincorporated area":
                suffix = "gemfr. Gebiet"
            elif r.get("ars") == "010545417035":
                suffix = "Kirchspiel"
            elif r.get("postal_code"):
                suffix = r["postal_code"]
            else:
                suffix = r.get("ars", "")

            if r["name"].endswith(")"):
                r["name"] = f"{r['name'][:-1]}, {suffix})"
            else:
                r["name"] = f"{r['name']} ({suffix})"
            r["slug"] = generate_slug(r["name"])

    # Pass 3: final safety guarantee to ensure absolute mathematical uniqueness
    seen_slugs: Dict[str, int] = {}
    for r in records:
        s = r["slug"]
        if s in seen_slugs:
            seen_slugs[s] += 1
            idx = seen_slugs[s]
            r["name"] = f"{r['name']} ({idx})"
            r["slug"] = f"{s}-{idx}"
        else:
            seen_slugs[s] = 1


def parse_float_coordinate(val: Any) -> Optional[float]:
    """Converts coordinate strings containing commas into Python floats."""
    if val is None:
        return None
    val_str = str(val).strip().replace(",", ".")
    if not val_str:
        return None
    try:
        return round(float(val_str), 6)
    except ValueError:
        return None


def parse_numeric(val: Any, is_int: bool = False) -> Optional[Any]:
    """Safely parse numeric values."""
    if val is None or val == "":
        return None
    if isinstance(val, (int, float)):
        return int(val) if is_int else round(float(val), 2)
    val_str = str(val).strip().replace(".", "").replace(",", ".")
    try:
        f = float(val_str)
        return int(f) if is_int else round(f, 2)
    except ValueError:
        return None


def format_plz(val: Any) -> str:
    """Formats postal code as a 5-digit string with leading zeros."""
    if val is None or val == "":
        return ""
    val_str = str(val).strip()
    if val_str.isdigit():
        return val_str.zfill(5)
    return val_str


def build_ars(r2: Any, r3: Any, r4: Any, r5: Any, r6: Any) -> str:
    """
    Constructs the 12-digit Official Regional Key (Amtlicher Regionalschlüssel - ARS):
    - State / Land (2 digits)
    - Administrative district / Regierungsbezirk (1 digit, e.g. '0')
    - District / Kreis (2 digits)
    - Municipality association / Gemeindeverband (4 digits)
    - Municipality / Gemeinde (3 digits)
    """
    land = str(r2 or "").strip().zfill(2)
    rb = str(r3 or "0").strip()
    kreis = str(r4 or "").strip().zfill(2)
    vb = str(r5 or "0000").strip().zfill(4)
    gem = str(r6 or "000").strip().zfill(3)
    return f"{land}{rb}{kreis}{vb}{gem}"


def build_ags(r2: Any, r3: Any, r4: Any, r6: Any) -> str:
    """
    Constructs the 8-digit Official Municipality Key (Amtlicher Gemeindeschlüssel - AGS):
    - State (2 digits) + Admin district (1 digit) + District (2 digits) + Municipality (3 digits)
    """
    land = str(r2 or "").strip().zfill(2)
    rb = str(r3 or "0").strip()
    kreis = str(r4 or "").strip().zfill(2)
    gem = str(r6 or "000").strip().zfill(3)
    return f"{land}{rb}{kreis}{gem}"


def strip_name_suffix(val: Optional[str]) -> Optional[str]:
    """Removes official title and status suffixes following a comma (e.g. ', Stadt', ', Hansestadt', ', M')."""
    if not val or "," not in val:
        return val
    return val.split(",", 1)[0].strip()


def find_data_sheet(workbook: openpyxl.Workbook) -> openpyxl.worksheet.worksheet.Worksheet:
    """Finds the worksheet containing the municipality data."""
    for name in workbook.sheetnames:
        if "Gemeinden" in name or "Onlineprodukt" in name:
            return workbook[name]
    # Fallback to second worksheet or active worksheet
    if len(workbook.sheetnames) > 1:
        return workbook.worksheets[1]
    return workbook.active


def extract_data(
    excel_path: str,
    inhabited_only: bool = False,
    include_extra: bool = False,
    strip_suffixes: bool = False,
) -> List[Dict[str, Any]]:
    """Reads the Excel file and transforms the data."""
    if not os.path.exists(excel_path):
        raise FileNotFoundError(f"File not found: {excel_path}")

    wb = openpyxl.load_workbook(excel_path, read_only=True, data_only=True)
    sheet = find_data_sheet(wb)

    current_land_name: Optional[str] = None
    current_kreis_name: Optional[str] = None
    current_verband_name: Optional[str] = None

    records: List[Dict[str, Any]] = []

    for idx, row in enumerate(sheet.iter_rows(values_only=True)):
        if idx < 6 or not row or not row[0]:
            continue

        satzart = str(row[0]).strip()

        # Record type 10: Federal state (Bundesland)
        if satzart == "10":
            current_land_name = str(row[7]).strip() if row[7] else None

        # Record type 40: District / independent city (Kreis / kreisfreie Stadt)
        elif satzart == "40":
            current_kreis_name = str(row[7]).strip() if row[7] else None

        # Record type 50: Municipality association / administration community (Gemeindeverband)
        elif satzart == "50":
            current_verband_name = str(row[7]).strip() if row[7] else None

        # Record type 60: Municipality / city (Gemeinde / Stadt)
        elif satzart == "60":
            name = str(row[7] or "").strip()
            einwohner = parse_numeric(row[9], is_int=True)

            # Optional: Filter only inhabited municipalities
            if inhabited_only and (einwohner is None or einwohner == 0):
                continue

            ars = build_ars(row[2], row[3], row[4], row[5], row[6])
            flaeche = parse_numeric(row[8], is_int=False)
            plz = format_plz(row[13])
            lon = parse_float_coordinate(row[14])
            lat = parse_float_coordinate(row[15])
            m_type = determine_type(row[1], name)
            district_name = strip_name_suffix(current_kreis_name) if strip_suffixes else current_kreis_name
            municipality_name = (strip_name_suffix(name) or name) if strip_suffixes else name
            slug = generate_slug(strip_name_suffix(name) or name)

            record: Dict[str, Any] = {
                "ars": ars,
                "state": current_land_name,
                "district": district_name,
                "name": municipality_name,
                "slug": slug,
                "type": m_type,
                "postal_code": plz,
                "country_code": "DE",
                "population": einwohner,
                "area_km2": flaeche,
                "latitude": lat,
                "longitude": lon,
            }

            if include_extra:
                record["ags"] = build_ags(row[2], row[3], row[4], row[6])
                record["municipality_association"] = current_verband_name
                record["designation"] = str(row[1]).strip() if row[1] else None
                record["population_male"] = parse_numeric(row[10], is_int=True)
                record["population_female"] = parse_numeric(row[11], is_int=True)
                record["density_per_km2"] = parse_numeric(row[12], is_int=False)
                record["travel_region"] = str(row[17]).strip() if len(row) > 17 and row[17] else None
                record["urbanization"] = str(row[19]).strip() if len(row) > 19 and row[19] else None

            records.append(record)

    wb.close()
    ensure_unique_slugs_and_names(records)
    return records


def export_json(data: List[Dict[str, Any]], filepath: str, indent: Optional[int] = 2) -> None:
    """Exports records as a clean UTF-8 JSON file."""
    os.makedirs(os.path.dirname(os.path.abspath(filepath)), exist_ok=True)
    with open(filepath, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=indent)


def export_csv(
    data: List[Dict[str, Any]],
    filepath: str,
    delimiter: str = ",",
    use_bom: bool = True,
) -> None:
    """
    Exports records as CSV.
    Defaults to UTF-8-BOM ('utf-8-sig') so that Microsoft Excel
    and other spreadsheet applications open special characters/umlauts properly.
    """
    if not data:
        return

    os.makedirs(os.path.dirname(os.path.abspath(filepath)), exist_ok=True)
    encoding = "utf-8-sig" if use_bom else "utf-8"
    fieldnames = list(data[0].keys())

    with open(filepath, "w", encoding=encoding, newline="") as f:
        writer = csv.DictWriter(
            f,
            fieldnames=fieldnames,
            delimiter=delimiter,
            quoting=csv.QUOTE_MINIMAL,
        )
        writer.writeheader()
        writer.writerows(data)


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Extracts Destatis GV3Q municipality data from Excel to JSON and CSV.",
        formatter_class=argparse.ArgumentDefaultsHelpFormatter,
    )
    parser.add_argument(
        "--input",
        "-i",
        default="AuszugGV3QAktuell.xlsx",
        help="Path to the Destatis Excel file",
    )
    parser.add_argument(
        "--output-json",
        "-j",
        default="german_municipalities.json",
        help="Output file path for the JSON export",
    )
    parser.add_argument(
        "--output-csv",
        "-c",
        default="german_municipalities.csv",
        help="Output file path for the CSV export",
    )
    parser.add_argument(
        "--csv-delimiter",
        default=",",
        help="Delimiter character for the CSV file (e.g., ',' or ';')",
    )
    parser.add_argument(
        "--inhabited-only",
        action="store_true",
        help="Export only inhabited municipalities (filters out uninhabited unincorporated areas)",
    )
    parser.add_argument(
        "--include-extra",
        action="store_true",
        help="Export additional attributes (AGS, gender breakdown, density, travel area, urbanization)",
    )
    parser.add_argument(
        "--compact-json",
        action="store_true",
        help="Compress JSON without indentation/whitespace (saves disk space)",
    )
    parser.add_argument(
        "--strip-suffixes",
        action="store_true",
        help="Strip official title and status suffixes following a comma (e.g. ', Stadt', ', Hansestadt', ', M')",
    )

    args = parser.parse_args()

    start_time = time.time()
    print(f"Reading '{args.input}'...")

    try:
        data = extract_data(
            excel_path=args.input,
            inhabited_only=args.inhabited_only,
            include_extra=args.include_extra,
            strip_suffixes=args.strip_suffixes,
        )
    except Exception as exc:
        print(f"Error while reading file: {exc}", file=sys.stderr)
        sys.exit(1)

    elapsed_read = time.time() - start_time
    print(f"Successfully loaded {len(data):,} municipalities in {elapsed_read:.2f}s.")

    # Export JSON
    indent = None if args.compact_json else 2
    export_json(data, args.output_json, indent=indent)
    json_size_mb = os.path.getsize(args.output_json) / (1024 * 1024)
    print(f"-> JSON saved: '{args.output_json}' ({json_size_mb:.2f} MB)")

    # Export CSV
    export_csv(data, args.output_csv, delimiter=args.csv_delimiter)
    csv_size_mb = os.path.getsize(args.output_csv) / (1024 * 1024)
    print(f"-> CSV saved:  '{args.output_csv}' ({csv_size_mb:.2f} MB)")

    total_time = time.time() - start_time
    print(f"Finished in {total_time:.2f} seconds.")


if __name__ == "__main__":
    main()
